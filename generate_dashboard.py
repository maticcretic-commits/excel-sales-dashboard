#!/usr/bin/env python3
"""Build sales-dashboard.xlsx — a formula-driven Excel sales dashboard.

Run: python3 generate_dashboard.py
Output: sales-dashboard.xlsx (in this directory)

The Dashboard sheet uses REAL Excel formulas (SUMIFS, COUNTIFS, AVERAGE,
XLOOKUP, EDATE) — nothing is hardcoded. Open the file in Excel / LibreOffice
and every KPI, summary table, and chart updates from the Data sheet.
"""

import random
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference

random.seed(42)

REGIONS = ["North", "South", "East", "West"]
REPS = {
    "North": ["Amit Sharma", "Priya Verma"],
    "South": ["Ravi Nair", "Divya Menon"],
    "East": ["Sourav Das", "Anita Roy"],
    "West": ["Vikram Mehta", "Kavita Patil"],
}
PRODUCTS = {
    "Laptop": ("Electronics", 55000),
    "Smartphone": ("Electronics", 22000),
    "Tablet": ("Electronics", 18000),
    "Monitor": ("Electronics", 14000),
    "Headphones": ("Accessories", 3500),
    "Keyboard": ("Accessories", 2500),
}

N_ROWS = 500
FIRST, LAST = date(2025, 1, 1), date(2025, 12, 31)
SPAN = (LAST - FIRST).days

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
KPI_LABEL_FONT = Font(bold=True, size=10, color="595959")
KPI_VALUE_FONT = Font(bold=True, size=14, color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RUPEE = '\u20B9#,##0'
RUPEE_DEC = '\u20B9#,##0.00'


def rand_date():
    from datetime import timedelta
    return FIRST + timedelta(days=random.randint(0, SPAN))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


wb = Workbook()

# ---------------------------------------------------------------- Data sheet
ws = wb.active
ws.title = "Data"
headers = ["OrderID", "Date", "Region", "SalesRep", "Product", "Category",
           "Units", "UnitPrice", "Revenue", "Cost", "Profit"]
ws.append(headers)
style_header(ws, 1, len(headers))

for i in range(N_ROWS):
    region = random.choice(REGIONS)
    product = random.choice(list(PRODUCTS))
    category, price = PRODUCTS[product]
    units = random.randint(1, 12)
    revenue = units * price
    cost = round(revenue * random.uniform(0.55, 0.75))
    ws.append([
        f"ORD-{1001 + i}",
        rand_date(),
        region,
        random.choice(REPS[region]),
        product,
        category,
        units,
        price,
        revenue,
        cost,
        revenue - cost,
    ])

for r in range(2, N_ROWS + 2):
    ws.cell(row=r, column=2).number_format = "DD-MMM-YYYY"
    for c in (8, 9, 10, 11):
        ws.cell(row=r, column=c).number_format = RUPEE
    ws.cell(row=r, column=7).alignment = Alignment(horizontal="center")

# Data validation: Region and Category dropdowns
dv_region = DataValidation(type="list", formula1='"North,South,East,West"',
                           allow_blank=False)
dv_region.error = "Pick a region from the list"
dv_region.prompt = "Choose region"
ws.add_data_validation(dv_region)
dv_region.add(f"C2:C{N_ROWS + 1}")

dv_cat = DataValidation(type="list",
                        formula1='"Electronics,Accessories"',
                        allow_blank=False)
ws.add_data_validation(dv_cat)
dv_cat.add(f"F2:F{N_ROWS + 1}")

# Conditional formatting: big deals green, thin-profit rows flagged
ws.conditional_formatting.add(
    f"I2:I{N_ROWS + 1}",
    CellIsRule(operator="greaterThan", formula=["100000"],
               fill=PatternFill("solid", fgColor="C6EFCE"),
               font=Font(color="006100")))
ws.conditional_formatting.add(
    f"K2:K{N_ROWS + 1}",
    CellIsRule(operator="lessThan", formula=["5000"],
               fill=PatternFill("solid", fgColor="FFC7CE"),
               font=Font(color="9C0006")))

widths = [12, 14, 10, 16, 14, 13, 8, 12, 14, 14, 14]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{N_ROWS + 1}"
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------ Dashboard sheet
db = wb.create_sheet("Dashboard")
db.sheet_view.showGridLines = False

db.merge_cells("A1:L1")
c = db["A1"]
c.value = "Sales Performance Dashboard — 2025"
c.font = TITLE_FONT
c.alignment = Alignment(horizontal="center", vertical="center")
db.row_dimensions[1].height = 30
db.merge_cells("A2:L2")
db["A2"] = "All figures below are live Excel formulas reading the Data sheet — change any row and the dashboard updates."
db["A2"].font = Font(italic=True, size=10, color="595959")
db["A2"].alignment = Alignment(horizontal="center")

# KPI cards (row 3 = label, row 4 = value)
kpis = [
    ("B", "Total Revenue", '=SUM(Data!I2:I501)', RUPEE),
    ("D", "Total Orders", "=COUNTA(Data!A2:A501)", '#,##0'),
    ("F", "Avg Order Value", '=AVERAGE(Data!I2:I501)', RUPEE),
    ("H", "Total Profit", '=SUM(Data!K2:K501)', RUPEE),
    ("J", "Top Product", '=XLOOKUP(MAX(K11:K16),K11:K16,J11:J16)', None),
]
for col, label, formula, fmt in kpis:
    db[f"{col}3"] = label
    db[f"{col}3"].font = KPI_LABEL_FONT
    db[f"{col}4"] = formula
    db[f"{col}4"].font = KPI_VALUE_FONT
    if fmt:
        db[f"{col}4"].number_format = fmt
    for r in (3, 4):
        db[f"{col}{r}"].border = BORDER
        db[f"{col}{r}"].alignment = Alignment(horizontal="center",
                                              vertical="center")
db.row_dimensions[4].height = 24

# Monthly summary table (A10:D22) — SUMIFS/COUNTIFS by month
db["A9"] = "Monthly trend (formula-driven)"
db["A9"].font = Font(bold=True, size=12, color="1F4E78")
for col, h in zip("ABCD", ["Month", "Revenue", "Orders", "Avg Order"]):
    db[f"{col}10"] = h
style_header(db, 10, 4)
for m in range(1, 13):
    r = 10 + m
    db.cell(row=r, column=1).value = date(2025, m, 1)
    db.cell(row=r, column=1).number_format = "mmm-yy"
    db.cell(row=r, column=2).value = (
        f'=SUMIFS(Data!I$2:I$501,Data!B$2:B$501,">="&A{r},'
        f'Data!B$2:B$501,"<"&EDATE(A{r},1))')
    db.cell(row=r, column=2).number_format = RUPEE
    db.cell(row=r, column=3).value = (
        f'=COUNTIFS(Data!B$2:B$501,">="&A{r},Data!B$2:B$501,"<"&EDATE(A{r},1))')
    db.cell(row=r, column=4).value = f'=IF(C{r}=0,0,B{r}/C{r})'
    db.cell(row=r, column=4).number_format = RUPEE
    for col in range(1, 5):
        db.cell(row=r, column=col).border = BORDER

# Region summary table (F10:H14) — SUMIFS per region
db["F9"] = "Revenue by region"
db["F9"].font = Font(bold=True, size=12, color="1F4E78")
for col, h in zip("FGH", ["Region", "Revenue", "Orders"]):
    db[f"{col}10"] = h
style_header(db, 10, 3)
for i, region in enumerate(REGIONS):
    r = 11 + i
    db[f"F{r}"] = region
    db[f"G{r}"] = f'=SUMIFS(Data!I$2:I$501,Data!C$2:C$501,F{r})'
    db[f"G{r}"].number_format = RUPEE
    db[f"H{r}"] = f'=COUNTIFS(Data!C$2:C$501,F{r})'
    for col in "FGH":
        db[f"{col}{r}"].border = BORDER

# Product summary table (J10:K16) — feeds the XLOOKUP top-product KPI
db["J9"] = "Revenue by product"
db["J9"].font = Font(bold=True, size=12, color="1F4E78")
for col, h in zip("JK", ["Product", "Revenue"]):
    db[f"{col}10"] = h
style_header(db, 10, 2)
for i, product in enumerate(PRODUCTS):
    r = 11 + i
    db[f"J{r}"] = product
    db[f"K{r}"] = f'=SUMIFS(Data!I$2:I$501,Data!E$2:E$501,J{r})'
    db[f"K{r}"].number_format = RUPEE
    for col in "JK":
        db[f"{col}{r}"].border = BORDER

for col, w in zip("ABCDEFGHIJKL",
                  [12, 16, 10, 16, 10, 16, 14, 10, 12, 16, 14, 12]):
    db.column_dimensions[col].width = w

# Native Excel charts -----------------------------------------------------
line = LineChart()
line.title = "Monthly Revenue Trend — 2025"
line.style = 10
line.y_axis.title = "Revenue (₹)"
line.x_axis.title = "Month"
data = Reference(db, min_col=2, min_row=10, max_row=22)
cats = Reference(db, min_col=1, min_row=11, max_row=22)
line.add_data(data, titles_from_data=True)
line.set_categories(cats)
line.height = 7.5
line.width = 15
db.add_chart(line, "A25")

bar = BarChart()
bar.type = "col"
bar.title = "Revenue by Region"
bar.style = 10
bar.y_axis.title = "Revenue (₹)"
data = Reference(db, min_col=7, min_row=10, max_row=14)
cats = Reference(db, min_col=6, min_row=11, max_row=14)
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.height = 7.5
bar.width = 13
db.add_chart(bar, "H25")

wb.save("sales-dashboard.xlsx")
print("Wrote sales-dashboard.xlsx")
